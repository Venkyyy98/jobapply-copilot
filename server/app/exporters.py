from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer


def _safe(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _format_cover_contact_line(line: str) -> str:
    segments = [seg.strip() for seg in line.split("|")]
    rendered: list[str] = []
    for seg in segments:
        if "linkedin.com/" in seg.lower():
            url = seg if seg.lower().startswith("http") else f"https://{seg}"
            rendered.append(f'<link href="{_safe(url)}" color="blue">{_safe(seg)}</link>')
        else:
            rendered.append(_safe(seg))
    return " | ".join(rendered)


def _split_cover_letter(text: str) -> dict[str, object]:
    lines = [line.rstrip() for line in text.splitlines()]
    nonempty = [line.strip() for line in lines if line.strip()]
    if len(nonempty) < 8:
        return {
            "name": "",
            "contact": "",
            "date": "",
            "recipient": [],
            "salutation": "",
            "body": [p.strip() for p in text.split("\n\n") if p.strip()],
            "closing": [],
        }

    name = nonempty[0]
    contact = nonempty[1] if len(nonempty) > 1 else ""
    date_line = nonempty[2] if len(nonempty) > 2 else ""
    salutation_index = next((idx for idx, line in enumerate(nonempty) if line.lower().startswith("dear ")), 6)
    recipient = nonempty[3:salutation_index]
    salutation = nonempty[salutation_index] if salutation_index < len(nonempty) else "Dear Hiring Manager,"
    closing_markers = {"sincerely,", "best regards,", "regards,"}
    closing_index = next(
        (idx for idx, line in enumerate(nonempty) if line.lower() in closing_markers),
        len(nonempty),
    )
    body = nonempty[salutation_index + 1 : closing_index]
    closing = nonempty[closing_index:] if closing_index < len(nonempty) else []
    return {
        "name": name,
        "contact": contact,
        "date": date_line,
        "recipient": recipient,
        "salutation": salutation,
        "body": body,
        "closing": closing,
    }


# Compatibility names retained for integrations importing the former exporters.
from .resume_export import export_pdf, export_docx, pdf_story, pdf_bytes, layout_status
from .resume_document import parse_resume

DEFAULT_RESUME_LAYOUT = {"body_font": 10.0, "docx_bullet_space_after": .65,
    "docx_paragraph_space_after": 1.2, "bullet_left_indent": 10,
    "bullet_first_line_indent": -6}
BALANCED_RESUME_LAYOUT = {"body_font": 10.0, "compact": True}
TIGHT_RESUME_LAYOUT = BALANCED_RESUME_LAYOUT


def _resume_lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines() if line.strip()]


def _resume_pdf_story(lines, layout):
    return pdf_story(parse_resume("\n".join(lines)), bool(layout.get("compact")))


def _render_resume_pdf_bytes(text, layout):
    return pdf_bytes(parse_resume(text), bool(layout.get("compact")))


def resume_layout_status(text: str) -> str:
    return layout_status(text)


def resume_text_to_pdf(text: str, output_path: Path) -> None:
    export_pdf(text, output_path)


def cover_letter_text_to_pdf(text: str, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=LETTER,
        leftMargin=0.72 * inch,
        rightMargin=0.72 * inch,
        topMargin=0.58 * inch,
        bottomMargin=0.58 * inch,
    )
    styles = getSampleStyleSheet()
    parsed = _split_cover_letter(text)
    name_style = ParagraphStyle(
        "CoverName",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=14.5,
        leading=16,
        alignment=1,
        spaceAfter=2,
    )
    contact_style = ParagraphStyle(
        "CoverContact",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8.8,
        leading=10.2,
        alignment=1,
        spaceAfter=6,
    )
    meta_style = ParagraphStyle(
        "CoverMeta",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.2,
        leading=12.4,
        spaceAfter=0.5,
    )
    salutation_style = ParagraphStyle(
        "CoverSalutation",
        parent=meta_style,
        spaceBefore=8,
        spaceAfter=6,
    )
    body = ParagraphStyle(
        "CoverBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.3,
        leading=13.2,
        spaceAfter=6,
    )
    closing_style = ParagraphStyle(
        "CoverClosing",
        parent=meta_style,
        spaceBefore=6,
        spaceAfter=1.5,
    )
    story = []
    if parsed["name"]:
        story.append(Paragraph(_safe(str(parsed["name"])), name_style))
    if parsed["contact"]:
        story.append(Paragraph(_format_cover_contact_line(str(parsed["contact"])), contact_style))
        story.append(HRFlowable(width="100%", thickness=0.55, color=colors.HexColor("#8A96A3"), spaceBefore=0, spaceAfter=10))
    if parsed["date"]:
        story.append(Paragraph(_safe(str(parsed["date"])), meta_style))
    for line in parsed["recipient"]:
        story.append(Paragraph(_safe(str(line)), meta_style))
    if parsed["salutation"]:
        story.append(Paragraph(_safe(str(parsed["salutation"])), salutation_style))
    for paragraph in parsed["body"]:
        story.append(Paragraph(_safe(str(paragraph)).replace("\n", "<br/>"), body))
    for line in parsed["closing"]:
        story.append(Paragraph(_safe(str(line)), closing_style))
    doc.build(story)


def text_to_pdf(text: str, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(output_path), pagesize=LETTER)
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontName = "Helvetica"
    body.fontSize = 10
    body.leading = 14

    story = []
    for paragraph in [p.strip() for p in text.split("\n\n") if p.strip()]:
        safe = paragraph.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe.replace("\n", "<br/>"), body))
        story.append(Spacer(1, 8))

    doc.build(story)


def text_to_docx(text: str, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    lines = [line.rstrip() for line in text.splitlines()]
    paragraph = None
    for line in lines:
        if not line.strip():
            paragraph = None
            continue
        if paragraph is None:
            paragraph = doc.add_paragraph(line.strip())
        else:
            paragraph.add_run("\n" + line.strip())
    doc.save(str(output_path))


def _set_page_margins(doc: Document, left: float, right: float, top: float, bottom: float) -> None:
    section = doc.sections[0]
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)


def _set_run_font(run, font_name: str, size: float | None = None) -> None:
    run.font.name = font_name
    if size is not None:
        run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font_name)


def _set_docx_default_font(doc: Document, font_name: str = "Arial") -> None:
    for style_name in ("Normal", "List Bullet"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = font_name
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn(f"w:{attr}"), font_name)
        try:
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        except Exception:
            pass


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), "Times New Roman")
    r_pr.append(r_fonts)
    r_pr.append(color)
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _docx_paragraph(
    doc: Document,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    size: float = 10.0,
    leading: float | None = None,
    align: WD_PARAGRAPH_ALIGNMENT | None = None,
    bullet: bool = False,
    layout: dict[str, float] = DEFAULT_RESUME_LAYOUT,
) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(
        layout["docx_bullet_space_after"] if bullet else layout["docx_paragraph_space_after"]
    )
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(leading or (size * 1.13))
    if bullet:
        p.paragraph_format.left_indent = Pt(layout["bullet_left_indent"])
        p.paragraph_format.first_line_indent = Pt(layout["bullet_first_line_indent"])
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    _set_run_font(run, "Times New Roman", size)


def resume_text_to_docx(text: str, output_path: Path) -> None:
    export_docx(text, output_path)


def cover_letter_text_to_docx(text: str, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _set_docx_default_font(doc, "Helvetica")
    _set_page_margins(doc, left=0.72, right=0.72, top=0.58, bottom=0.58)
    parsed = _split_cover_letter(text)

    if parsed["name"]:
        _docx_paragraph(
            doc,
            str(parsed["name"]),
            bold=True,
            size=14.5,
            align=WD_PARAGRAPH_ALIGNMENT.CENTER,
        )

    if parsed["contact"]:
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.paragraph_format.space_after = Pt(8)
        segments = [seg.strip() for seg in str(parsed["contact"]).split("|")]
        first = True
        for seg in segments:
            if not first:
                sep = p.add_run("  |  ")
                _set_run_font(sep, "Helvetica", 8.8)
            if "linkedin.com/" in seg.lower():
                url = seg if seg.lower().startswith("http") else f"https://{seg}"
                _add_hyperlink(p, seg, url)
            else:
                run = p.add_run(seg)
                _set_run_font(run, "Helvetica", 8.8)
            first = False

        border_p = doc.add_paragraph()
        border_p.paragraph_format.space_after = Pt(8)
        p_pr = border_p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "8A96A3")
        p_bdr.append(bottom)
        p_pr.append(p_bdr)

    if parsed["date"]:
        _docx_paragraph(doc, str(parsed["date"]), size=10.2, align=WD_PARAGRAPH_ALIGNMENT.LEFT)
    for line in parsed["recipient"]:
        _docx_paragraph(doc, str(line), size=10.2, align=WD_PARAGRAPH_ALIGNMENT.LEFT)
    if parsed["salutation"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(str(parsed["salutation"]))
        _set_run_font(run, "Helvetica", 10.3)
    for para in parsed["body"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.08
        run = p.add_run(str(para))
        _set_run_font(run, "Helvetica", 10.3)
    for idx, line in enumerate(parsed["closing"]):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5 if idx == 0 else 0)
        p.paragraph_format.space_after = Pt(1.5)
        run = p.add_run(str(line))
        _set_run_font(run, "Helvetica", 10.3)

    doc.save(str(output_path))
